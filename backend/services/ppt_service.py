from pptx import Presentation
from pptx.util import Inches, Pt
import os

REPORTS_DIR = "reports"
os.makedirs(REPORTS_DIR, exist_ok=True)


def generate_ppt_report(dataset_name: str, summary: str, kpis: list, recommendations: list) -> str:
    filename = os.path.join(REPORTS_DIR, f"report_{dataset_name}.pptx")
    prs = Presentation()

    # Slide 1 — Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"BI Report — {dataset_name}"
    slide.placeholders[1].text = "Autonomous Business Intelligence Platform"

    # Slide 2 — Executive Summary
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive Summary"
    slide.placeholders[1].text = summary

    # Slide 3 — KPIs
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Key Performance Indicators"
    body = slide.placeholders[1].text_frame
    body.text = f"{kpis[0]['measure']}: Total {kpis[0]['total']}"
    for k in kpis[1:]:
        p = body.add_paragraph()
        p.text = f"{k['measure']}: Total {k['total']}, Avg {k['average']}"

    # Slide 4 — Recommendations
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Recommendations"
    body = slide.placeholders[1].text_frame
    body.text = recommendations[0] if recommendations else "No recommendations available"
    for rec in recommendations[1:]:
        p = body.add_paragraph()
        p.text = rec

    prs.save(filename)
    return filename